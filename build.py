# Renders content.py to index.html (styled like pa-us-call-coverage) and a Word doc.
# Run: python build.py
import html
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

import content as C

HERE = Path(__file__).parent
DOCX_NAME = "Moghees Siddique - Roles and Responsibilities (October 2026).docx"
e = html.escape

CSS = """
:root {
  --bg: #f6f8f3; --surface: #ffffff; --fg: #16190f; --muted: #5f6655; --line: #dde3d4;
  --accent: #5f9e1f; --green: #7ec837; --blue: #3b5bdb; --amber: #b8860b; --amber-bg: #fdf6e3;
  --font-body: "Inter", system-ui, -apple-system, "Segoe UI", sans-serif;
}
@media (prefers-color-scheme: dark) {
  :root:not([data-theme="light"]) {
    --bg: #11140d; --surface: #1a1e15; --fg: #eef2e6; --muted: #a3ab96; --line: #2d3324;
    --accent: #9bdc5a; --blue: #7c93f0; --amber: #e0b54a; --amber-bg: #2a2412; color-scheme: dark;
  }
}
:root[data-theme="dark"] {
  --bg: #11140d; --surface: #1a1e15; --fg: #eef2e6; --muted: #a3ab96; --line: #2d3324;
  --accent: #9bdc5a; --blue: #7c93f0; --amber: #e0b54a; --amber-bg: #2a2412; color-scheme: dark;
}
* { box-sizing: border-box; }
html { scroll-behavior: smooth; }
body { margin: 0; background: var(--bg); color: var(--fg); font-family: var(--font-body); font-size: 15px; line-height: 1.55; padding: 0 16px; }
main { max-width: 900px; margin: 0 auto; padding-block: 40px 64px; display: flex; flex-direction: column; gap: 40px; }
h1, h2, h3 { text-wrap: balance; margin: 0; }
h1 { font-size: clamp(28px, 5vw, 40px); font-weight: 800; letter-spacing: -0.02em; line-height: 1.1; }
h2 { font-size: 24px; font-weight: 700; letter-spacing: -0.01em; }
h3 { font-size: 16px; font-weight: 700; margin-top: 6px; }
p { margin: 0; max-width: 68ch; }
.eyebrow { font-size: 12px; font-weight: 600; letter-spacing: 0.08em; text-transform: uppercase; color: var(--accent); }
.top { background: #0b0b0b; color: #fff; margin-inline: -16px; padding: 28px 16px 36px; border-bottom: 4px solid #7ec837; }
.top-in { max-width: 900px; margin: 0 auto; display: flex; flex-direction: column; gap: 10px; }
.top .logo { width: 76px; height: auto; margin-bottom: 6px; }
.top h1 { color: #fff; }
.top h1 .hl { color: #7ec837; }
.top .sub { color: #c9cfbf; font-size: 15px; }
section.part { display: flex; flex-direction: column; gap: 14px; border-top: 3px solid var(--fg); padding-top: 22px; scroll-margin-top: 16px; }
section.part > .eyebrow { margin-bottom: -8px; }

.stats { display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 12px; }
.stat { font: inherit; text-align: left; color: inherit; cursor: pointer; background: var(--surface); border: 1px solid var(--line); border-radius: 10px; padding: 16px; display: flex; flex-direction: column; gap: 4px; transition: border-color .15s, transform .15s; }
.stat:hover { border-color: var(--green); transform: translateY(-1px); }
.stat:focus-visible { outline: 2px solid var(--green); outline-offset: 2px; }
.stat .big { font-size: 32px; font-weight: 800; letter-spacing: -0.02em; color: var(--accent); line-height: 1.1; }
.stat .lab { font-size: 14px; color: var(--muted); }
.stat .more { font-size: 12px; font-weight: 600; color: var(--blue); margin-top: 6px; }
dialog.pop { width: min(520px, calc(100% - 32px)); max-height: min(80vh, 680px); padding: 0; border: 1px solid var(--line); border-radius: 14px; background: var(--surface); color: var(--fg); box-shadow: 0 20px 60px rgba(0,0,0,.25); }
dialog.pop::backdrop { background: rgba(11,11,11,.55); }
.pop-head { position: sticky; top: 0; display: flex; justify-content: space-between; align-items: start; gap: 12px; padding: 16px 18px 12px; background: var(--surface); border-bottom: 1px solid var(--line); border-top: 4px solid var(--green); }
.pop-head h3 { margin: 0; font-size: 17px; }
.pop-head .n { color: var(--accent); }
.pop-x { font: inherit; font-size: 22px; line-height: 1; background: none; border: 0; color: var(--muted); cursor: pointer; padding: 2px 6px; border-radius: 6px; }
.pop-x:hover { color: var(--fg); background: var(--bg); }
.pop-x:focus { outline: none; }
.pop-x:focus-visible { outline: 2px solid var(--green); }
.pop ul { list-style: none; margin: 0; padding: 6px 18px 16px; }
.pop li { padding: 9px 0; border-bottom: 1px solid var(--line); font-size: 15px; }
.pop li:last-child { border-bottom: 0; }
.pop li b { display: block; font-size: 12px; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: var(--accent); margin-bottom: 2px; }
.pop li small { color: var(--amber); font-weight: 600; margin-left: 6px; }
.cc { display: flex; flex-wrap: wrap; align-items: center; justify-content: space-between; gap: 12px; background: var(--surface); border: 1px solid var(--line); border-left: 4px solid var(--green); border-radius: 10px; padding: 14px 16px; }
.cc p { font-size: 14px; color: var(--muted); }
.cc p b { color: var(--fg); }
.cc a.btn { background: #0b0b0b; color: #fff; text-decoration: none; font-weight: 700; font-size: 14px; padding: 10px 16px; border-radius: 8px; white-space: nowrap; }
.cc a.btn:hover { background: #2a2f22; }
.card .where a { color: var(--blue); text-underline-offset: 3px; }

.toc { background: var(--surface); border: 1px solid var(--line); border-radius: 10px; padding: 16px 18px; }
.toc h2 { font-size: 15px; margin-bottom: 4px; }
.toc .hint { font-size: 13px; color: var(--muted); margin-bottom: 10px; }
.toc ol { margin: 0; padding: 0; list-style: none; display: grid; grid-template-columns: repeat(auto-fit, minmax(260px, 1fr)); gap: 4px 20px; }
.toc a { display: grid; grid-template-columns: 92px minmax(0, 1fr); gap: 10px; align-items: baseline; padding: 4px 0; color: var(--fg); text-decoration: none; font-weight: 500; }
.toc a span { font-size: 11px; font-weight: 600; letter-spacing: 0.06em; text-transform: uppercase; color: var(--accent); }
.toc a em { font-style: normal; color: var(--blue); text-decoration: underline; text-decoration-color: color-mix(in srgb, var(--blue) 40%, transparent); text-underline-offset: 3px; }
.toc a:hover em { text-decoration-color: var(--blue); }

ol.flow { margin: 0; padding: 0; list-style: none; counter-reset: s; display: flex; flex-direction: column; }
ol.flow li { counter-increment: s; position: relative; padding: 0 0 18px 48px; }
ol.flow li::before { content: counter(s); position: absolute; left: 0; top: 0; width: 32px; height: 32px; border-radius: 50%; background: var(--green); color: #0b0b0b; font-weight: 800; display: grid; place-items: center; font-size: 14px; }
ol.flow li::after { content: ""; position: absolute; left: 15px; top: 34px; bottom: 2px; width: 2px; background: var(--line); }
ol.flow li:last-child::after { display: none; }
ol.flow b { display: block; font-size: 16px; margin-bottom: 2px; padding-top: 4px; }

ul.list { margin: 0; padding: 0; list-style: none; background: var(--surface); border: 1px solid var(--line); border-radius: 10px; }
ul.list li { display: grid; grid-template-columns: 210px minmax(0, 1fr); gap: 14px; padding: 12px 16px; border-bottom: 1px solid var(--line); }
ul.list li:last-child { border-bottom: 0; }
ul.list b { font-weight: 600; }

.table-wrap { overflow-x: auto; background: var(--surface); border: 1px solid var(--line); border-radius: 10px; }
table { width: 100%; border-collapse: collapse; }
th, td { text-align: left; padding: 11px 14px; border-bottom: 1px solid var(--line); vertical-align: top; }
tr:last-child td { border-bottom: 0; }
th { font-size: 12px; font-weight: 600; letter-spacing: 0.06em; text-transform: uppercase; color: var(--muted); }
td:first-child { font-weight: 600; white-space: nowrap; }

.group { display: flex; flex-direction: column; gap: 10px; margin-top: 8px; }
.group > h3 { font-size: 13px; letter-spacing: 0.08em; text-transform: uppercase; color: var(--muted); }
.cards { display: grid; grid-template-columns: repeat(auto-fill, minmax(min(100%, 400px), 1fr)); gap: 12px; }
.card { background: var(--surface); border: 1px solid var(--line); border-top: 4px solid var(--green); border-radius: 10px; padding: 14px 16px; display: flex; flex-direction: column; gap: 8px; min-width: 0; }
.card.testing { border-top-color: var(--amber); }
.card header { display: flex; justify-content: space-between; align-items: start; gap: 10px; }
.card h4 { margin: 0; font-size: 17px; }
.card .where { font-size: 13px; color: var(--muted); overflow-wrap: anywhere; }
.pill { font-size: 11px; font-weight: 700; letter-spacing: 0.05em; text-transform: uppercase; padding: 3px 8px; border-radius: 99px; background: color-mix(in srgb, var(--green) 22%, transparent); color: var(--accent); white-space: nowrap; }
.testing .pill { background: var(--amber-bg); color: var(--amber); }
.card dl { margin: 0; display: grid; grid-template-columns: 92px minmax(0, 1fr); gap: 6px 10px; font-size: 14px; }
.card dt { font-size: 11px; font-weight: 600; letter-spacing: 0.06em; text-transform: uppercase; color: var(--muted); padding-top: 2px; }
.card dd { margin: 0; }

footer { color: var(--muted); font-size: 13px; border-top: 1px solid var(--line); padding-top: 16px; }

@media (max-width: 640px) {
  main { padding-block: 24px 48px; gap: 32px; }
  body { font-size: 16px; }
  .top { padding: 22px 16px 28px; } .top .logo { width: 64px; }
  h2 { font-size: 21px; }
  ul.list li { grid-template-columns: 1fr; gap: 2px; }
  .toc a { grid-template-columns: 80px minmax(0, 1fr); }
  table, thead, tbody, tr, td { display: block; }
  thead { display: none; }
  tr { padding: 10px 14px; border-bottom: 1px solid var(--line); }
  tr:last-child { border-bottom: 0; }
  td { padding: 2px 0; border: 0; white-space: normal !important; }
  td:not(:first-child) { color: var(--muted); }
  .card dl { grid-template-columns: 1fr; gap: 0; }
  .card dd { margin-bottom: 6px; }
}
@media print {
  body { background: #fff; }
  .toc { display: none; }
  .card, ul.list li, tr { break-inside: avoid; }
  section.part { break-before: auto; }
}
"""


def h_blocks(blocks):
    out = []
    for b in blocks:
        kind = b[0]
        if kind == "p":
            out.append(f"<p>{e(b[1])}</p>")
        elif kind == "h3":
            out.append(f"<h3>{e(b[1])}</h3>")
        elif kind == "bullets":
            out.append('<ul class="list">' + "".join(f"<li><b>{e(a)}</b><span>{e(t)}</span></li>" for a, t in b[1]) + "</ul>")
        elif kind == "flow":
            out.append('<ol class="flow">' + "".join(f"<li><b>{e(a)}</b>{e(t)}</li>" for a, t in b[1]) + "</ol>")
        elif kind == "table":
            head = "".join(f"<th>{e(x)}</th>" for x in b[1])
            rows = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in b[2])
            out.append(f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>')
        elif kind == "systems":
            for g in b[1]:
                cards = []
                for s in g["items"]:
                    cls = "card testing" if s["status"] != "Live" else "card"
                    cards.append(
                        f'<article class="{cls}"><header><div><h4>{e(s["name"])}</h4><div class="where">{where_html(s["where"])}</div></div>'
                        f'<span class="pill">{e(s["status"])}</span></header><p>{e(s["what"])}</p><dl>'
                        f'<dt>Who uses it</dt><dd>{e(s["who"])}</dd><dt>When</dt><dd>{e(s["when"])}</dd>'
                        f'<dt>If it stops</dt><dd>{e(s["stop"])}</dd></dl></article>')
                out.append(f'<div class="group"><h3>{e(g["name"])}</h3><div class="cards">{"".join(cards)}</div></div>')
    return "\n".join(out)


SCRIPT = """<script>
document.querySelectorAll(".stat").forEach(function (b) {
  b.addEventListener("click", function () { document.getElementById(b.dataset.pop).showModal(); });
});
document.querySelectorAll("dialog.pop").forEach(function (d) {
  d.querySelector(".pop-x").addEventListener("click", function () { d.close(); });
  d.addEventListener("click", function (ev) { if (ev.target === d) d.close(); });
});
</script>"""


def where_html(w):
    # Only the physiqueacademy.com addresses become links; each one has its own login.
    if "physiqueacademy.com" in w and " " not in w:
        return f'<a href="https://{e(w)}" target="_blank" rel="noopener">{e(w)}</a>'
    return e(w)


def find_block(kind, section_id=None):
    for s in C.SECTIONS:
        if section_id and s["id"] != section_id:
            continue
        for b in s["blocks"]:
            if b[0] == kind:
                return b


def stat_items(src, html_out=False):
    if isinstance(src, list):
        return [e(x) if html_out else x for x in src]
    if src == "systems":
        out = []
        for g in find_block("systems")[1]:
            for it in g["items"]:
                tag = "" if it["status"] == "Live" else (f' <small>{e(it["status"])}</small>' if html_out else f' ({it["status"]})')
                parts = f' ({it["parts"]} parts)' if it.get("parts") else ""
                out.append((e(it["name"]) if html_out else it["name"]) + parts + tag)
        return out
    if src == "platforms":
        return [e(r[0]) if html_out else r[0] for r in find_block("table", "platforms")[2]]
    if src == "reports":
        return [(f"<b>{e(r[0])}</b>{e(r[1])}" if html_out else f"{r[0]}: {r[1]}") for r in find_block("table", "rhythm")[2]]
    return []


def build_html():
    toc = "".join(f'<li><a href="#{s["id"]}"><span>{e(s["eyebrow"])}</span><em>{e(s["title"].split(":")[0])}</em></a></li>' for s in C.SECTIONS)
    stats = "".join(
        f'<button class="stat" type="button" data-pop="pop{i}"><span class="big">{e(a)}</span><span class="lab">{e(t)}</span><span class="more">Show list</span></button>'
        for i, (a, t, src) in enumerate(C.STATS))
    pops = "".join(
        f'<dialog class="pop" id="pop{i}" aria-labelledby="pop{i}h"><div class="pop-head"><h3 id="pop{i}h"><span class="n">{e(a)}</span> {e(t)}</h3>'
        f'<button class="pop-x" type="button" aria-label="Close">&times;</button></div>'
        f'<ul>{"".join(f"<li>{x}</li>" for x in stat_items(src, html_out=True))}</ul></dialog>'
        for i, (a, t, src) in enumerate(C.STATS))
    parts = "".join(
        f'<section class="part" id="{s["id"]}"><div class="eyebrow">{e(s["eyebrow"])}</div><h2>{e(s["title"])}</h2>{h_blocks(s["blocks"])}</section>'
        for s in C.SECTIONS)
    page = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="robots" content="noindex, nofollow">
<title>Roles and Responsibilities</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap">
<style>{CSS}</style>
</head>
<body>
<header class="top">
  <div class="top-in">
    <img class="logo" src="pa-logo-white.png" alt="Physique Academy" width="120" height="72">
    <h1><span class="hl">{e(C.TITLE)}:</span> {e(C.SUBTITLE[0].lower() + C.SUBTITLE[1:])}</h1>
    <div class="sub">{e(C.UPDATED)}</div>
  </div>
</header>
<main>
  <div class="cc"><p><b>Want to open the dashboards?</b> Every one is linked from your Command Center.</p><a class="btn" href="{C.COMMAND_CENTER}" target="_blank" rel="noopener">Open Command Center</a></div>
  <div class="stats">{stats}</div>
  <nav class="toc" aria-label="On this page"><h2>On this page</h2><p class="hint">Click a section to jump straight to it.</p><ol>{toc}</ol></nav>
  {parts}
  {pops}
  <footer>{e(C.TITLE)} &middot; {e(C.UPDATED)}. No passwords or keys are kept on this page.</footer>
</main>
{SCRIPT}
</body>
</html>
"""
    (HERE / "index.html").write_text(page, encoding="utf-8")


# ---------- Word ----------
GREEN = RGBColor(0x5F, 0x9E, 0x1F)
MUTED = RGBColor(0x5F, 0x66, 0x55)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def d_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
        shade(c, "E9F3DD")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(9.5)
            if i == 0: r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def d_blocks(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "h3":
            doc.add_heading(b[1], level=3)
        elif kind == "bullets":
            for a, t in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(a + ": ").bold = True
                p.add_run(t)
        elif kind == "flow":
            for i, (a, t) in enumerate(b[1], 1):
                p = doc.add_paragraph(style="List Number")
                p.add_run(a + ". ").bold = True
                p.add_run(t)
        elif kind == "table":
            d_table(doc, b[1], b[2], [3.5, 8.5, 5] if len(b[1]) == 3 else [4.5, 12.5])
        elif kind == "systems":
            for g in b[1]:
                doc.add_heading(g["name"], level=3)
                for s in g["items"]:
                    t = doc.add_table(rows=0, cols=2)
                    t.style = "Table Grid"
                    head = t.add_row().cells
                    m = head[0].merge(head[1])
                    m.text = ""
                    r = m.paragraphs[0].add_run(s["name"]); r.bold = True; r.font.size = Pt(11)
                    r2 = m.paragraphs[0].add_run(f'   {s["where"]}   [{s["status"]}]'); r2.font.size = Pt(9); r2.font.color.rgb = MUTED
                    shade(m, "E9F3DD" if s["status"] == "Live" else "FDF0D0")
                    for label, key in (("What it does", "what"), ("Who uses it", "who"), ("When", "when"), ("If it stops", "stop")):
                        c = t.add_row().cells
                        c[0].text = ""; c[1].text = ""
                        rr = c[0].paragraphs[0].add_run(label); rr.bold = True; rr.font.size = Pt(9.5)
                        c[1].paragraphs[0].add_run(s[key]).font.size = Pt(9.5)
                    for row in t.rows:
                        row.cells[0].width = Cm(3.5); row.cells[1].width = Cm(13.5)
                    doc.add_paragraph()


def build_docx():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        hs = doc.styles[f"Heading {lvl}"]; hs.font.name = "Calibri"; hs.font.size = Pt(size); hs.font.color.rgb = RGBColor(0x16, 0x19, 0x0F)

    logo = HERE / "pa-logo-dark.png"
    if logo.exists():
        doc.add_picture(str(logo), width=Cm(3))
    p = doc.add_paragraph(); r = p.add_run(C.TITLE); r.bold = True; r.font.size = Pt(24)
    p = doc.add_paragraph(); r = p.add_run(C.SUBTITLE); r.font.size = Pt(14); r.font.color.rgb = GREEN
    p = doc.add_paragraph(); r = p.add_run(C.UPDATED); r.font.color.rgb = MUTED

    d_table(doc, ["At a glance", ""], [[a, t + ":\n" + "\n".join("- " + x for x in stat_items(src))] for a, t, src in C.STATS], [3, 14])

    for s in C.SECTIONS:
        title = s["title"] if s["eyebrow"] == "Start here" else f'{s["eyebrow"].replace("Section ", "")}. {s["title"]}'
        doc.add_heading(title, level=1)
        d_blocks(doc, s["blocks"])
    out = HERE / DOCX_NAME
    doc.save(out)
    return out


if __name__ == "__main__":
    build_html()
    print("index.html written")
    print(build_docx(), "written")
